#!/usr/bin/env python3
"""Finalize the MnROAD Cell 22 observation-lineage audit.

The raw user-drawn centrelines are preserved.  A separate canonical axial
occupancy representation collapses same-page, same-date overlapping redraws
using the already-frozen temporal matching tolerances.  The 2019-10-23 maps
precede the documented 2019-10-24 sealing intervention and remain candidate
pre-intervention observations; no transition may cross that intervention.
"""

from __future__ import annotations

import argparse
import csv
import importlib.util
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ELIGIBLE_DATES = (
    "2017-03-20", "2018-03-14", "2018-04-27",
    "2018-08-23", "2019-03-22", "2019-04-16", "2019-10-23",
)
INTERVENTION_DATE = "2019-10-24"
LATERAL_TOLERANCE_PX = 6.0
LONGITUDINAL_GAP_TOLERANCE_FT = 0.5


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def atomic_json(path: Path, payload: dict) -> None:
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2) + "\n")
    temporary.replace(path)


def union_length(intervals: list[tuple[float, float]]) -> tuple[float, list[list[float]]]:
    merged: list[list[float]] = []
    for left, right in sorted(intervals):
        left, right = max(0.0, left), min(50.0, right)
        if right <= left:
            continue
        if not merged or left > merged[-1][1]:
            merged.append([left, right])
        else:
            merged[-1][1] = max(merged[-1][1], right)
    return sum(right - left for left, right in merged), merged


def groups(descriptors: list[dict]) -> list[list[dict]]:
    parent = list(range(len(descriptors)))

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def join(first: int, second: int) -> None:
        first, second = find(first), find(second)
        if first != second:
            parent[second] = first

    for first, left in enumerate(descriptors):
        for second in range(first):
            right = descriptors[second]
            gap = max(
                0.0,
                left["x_min_ft"] - right["x_max_ft"],
                right["x_min_ft"] - left["x_max_ft"],
            )
            lateral = abs(left["y_mean_px"] - right["y_mean_px"])
            if gap <= LONGITUDINAL_GAP_TOLERANCE_FT and lateral <= LATERAL_TOLERANCE_PX:
                join(first, second)
    result: dict[int, list[dict]] = defaultdict(list)
    for index, descriptor in enumerate(descriptors):
        result[find(index)].append(descriptor)
    return list(result.values())


def excel_totals(path: Path) -> dict[str, float]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Export Worksheet"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    result = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if row.get("SECTION_ID") != "22-D-2016" or row.get("LANE") != "Driving":
            continue
        date = row["DAY"].date().isoformat()
        if date in ELIGIBLE_DATES:
            result[date] = sum(float(row.get(key) or 0) for key in (
                "LONG_NWP_L_L", "LONG_NWP_L_M", "LONG_NWP_L_H",
            ))
    workbook.close()
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--packet", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--annotator", type=Path, required=True)
    parser.add_argument("--candidate-builder", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--lock", action="store_true")
    args = parser.parse_args()

    packet = args.packet.resolve()
    output = args.output.resolve()
    annotator = load_module("mnroad_annotator", args.annotator.resolve())
    helper = load_module("mnroad_candidate_helper", args.candidate_builder.resolve())
    private = json.loads((packet / "private_repeat_key.json").read_text())["tasks"]
    records = {
        task_id: json.loads((packet / "records" / f"{task_id}.geojson").read_text())
        for task_id in private
    }
    if not all(record.get("phase1_locked") for record in records.values()):
        raise RuntimeError("all phase-1 records must be locked")
    if args.lock and any(record.get("phase2_locked") for record in records.values()):
        raise RuntimeError("phase 2 is already partly or fully locked")

    primary = {
        (row["survey_date"], row["page"]): task_id
        for task_id, row in private.items() if not row["is_repeat"]
    }
    canonical = []
    date_totals = Counter()
    for date in ELIGIBLE_DATES:
        for page in range(1, 11):
            task_id = primary[(date, page)]
            task = private[task_id]
            record = records[task_id]
            width, height = annotator.image_size(packet / "images" / task["image"])
            descriptors = [
                helper.descriptor(
                    feature, date, width, height,
                    annotator.landscape_rotation_degrees(date),
                )
                for feature in annotator._target_features(record)
            ]
            for cluster_index, cluster in enumerate(groups(descriptors), 1):
                length, intervals = union_length([
                    (item["x_min_ft"], item["x_max_ft"]) for item in cluster
                ])
                source_segments = sorted(item["segment_id"] for item in cluster)
                source_lineages = sorted({
                    item["feature"]["properties"].get("lineage_id", "")
                    for item in cluster
                    if item["feature"]["properties"].get("lineage_id")
                })
                relations = sorted({
                    item["feature"]["properties"].get("relation_type", "unresolved")
                    for item in cluster
                })
                row = {
                    "canonical_id": f"CAN-{task_id}-{cluster_index:02d}",
                    "task_id": task_id,
                    "survey_date": date,
                    "page_start_ft": record["page_start_ft"],
                    "page_end_ft": record["page_end_ft"],
                    "axial_intervals_ft": intervals,
                    "canonical_axial_length_ft": length,
                    "lateral_mean_px": sum(item["y_mean_px"] for item in cluster) / len(cluster),
                    "source_segment_ids": source_segments,
                    "source_lineage_ids": source_lineages,
                    "source_relation_types": relations,
                    "source_segment_count": len(cluster),
                    "canonical_relation": "redraw_union" if len(cluster) > 1 else relations[0],
                }
                canonical.append(row)
                date_totals[date] += length

    excel = excel_totals(args.workbook.resolve())
    reconciliation = []
    for date in ELIGIBLE_DATES:
        map_total = date_totals[date]
        excel_total = excel[date]
        tolerance = max(5.0, 0.10 * excel_total)
        reconciliation.append({
            "survey_date": date,
            "excel_nwp_ft": excel_total,
            "canonical_axial_ft": map_total,
            "difference_ft": map_total - excel_total,
            "tolerance_ft": tolerance,
            "passes": abs(map_total - excel_total) <= tolerance,
        })

    low_confidence = []
    for task_id, record in records.items():
        for feature in record["features"]:
            properties = feature["properties"]
            if properties.get("mapping_confidence") == "low":
                properties["phase2_review_status"] = "retained_explicit_unresolved"
                properties["change_type"] = "unresolvable"
                low_confidence.append(properties["segment_id"])
        record.pop("trajectory_censor_reason", None)
        if record["survey_date"] in ELIGIBLE_DATES:
            record["trajectory_eligibility"] = "natural_observation_candidate"

    timestamp = datetime.now(timezone.utc).isoformat()
    if args.lock:
        for task_id, record in records.items():
            record["phase2_locked"] = True
            record["phase2_locked_at_utc"] = timestamp
            record["phase2_lock_status"] = "locked_with_explicit_unresolved_and_intervention_boundary"
            atomic_json(packet / "records" / f"{task_id}.geojson", record)

    m0042 = [row for row in canonical if row["task_id"] == "M0042"]
    phase2_locked = args.lock or all(record.get("phase2_locked") for record in records.values())
    result = {
        "status": "PHASE2_LOCKED__QUALIFICATION_STILL_NO_GO" if phase2_locked else "PHASE2_READY_TO_LOCK__QUALIFICATION_STILL_NO_GO",
        "raw_geometry_preserved": True,
        "canonical_method": {
            "representation": "same-page same-date axial interval union",
            "lateral_tolerance_px": LATERAL_TOLERANCE_PX,
            "longitudinal_gap_tolerance_ft": LONGITUDINAL_GAP_TOLERANCE_FT,
            "threshold_provenance": "reuses frozen phase-2 temporal matching tolerance; not tuned to Excel",
        },
        "intervention_boundary": {
            "date": INTERVENTION_DATE,
            "decision": "retain 2019-10-23 as pre-intervention candidate; forbid transitions across intervention",
            "last_eligible_pre_intervention_observation": "2019-10-23",
        },
        "canonical_record_count": len(canonical),
        "m0042_canonical_records": m0042,
        "low_confidence_segments_retained_unresolved": sorted(low_confidence),
        "reconciliation": reconciliation,
        "passed_dates": sum(row["passes"] for row in reconciliation),
        "eligible_dates": len(reconciliation),
        "qualification": "NO_GO_MODEL_READY_TRAJECTORY",
        "reason": "2018 onset/growth scalar reconciliation still fails and only one section is available",
    }
    output.mkdir(parents=True, exist_ok=True)
    atomic_json(output / "canonical_observation_state.json", {
        "type": "MnROADCanonicalAxialObservationCollection",
        "records": canonical,
    })
    with (output / "canonical_observation_state.csv").open("w", newline="") as handle:
        fields = [
            "canonical_id", "task_id", "survey_date", "page_start_ft", "page_end_ft",
            "canonical_axial_length_ft", "lateral_mean_px", "source_segment_count",
            "source_segment_ids", "source_lineage_ids", "source_relation_types", "canonical_relation",
        ]
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in canonical:
            writer.writerow({key: row.get(key) for key in fields})
    atomic_json(output / "phase2_finalization_report.json", result)
    markdown = [
        "# MnROAD Cell 22 phase-2 finalisation",
        "",
        f"Status: `{result['status']}`",
        "",
        "## Decisions",
        "",
        "- Raw user-drawn geometry is preserved.",
        f"- Revised M0042 annotation is represented by {len(m0042)} canonical axial component(s) "
        f"totalling {sum(row['canonical_axial_length_ft'] for row in m0042):.1f} ft; "
        f"{len(records['M0042']['features'])} raw source segment(s) remain unchanged.",
        f"- {len(low_confidence)} low-confidence lineage mapping(s) are retained explicitly as `unresolvable`.",
        "- 2019-10-23 is retained as a pre-intervention candidate; no transition may cross the documented 2019-10-24 sealing intervention.",
        "",
        "## Excel reconciliation after canonical redraw union",
        "",
        "| Date | Excel NWP (ft) | Canonical axial occupancy (ft) | Difference (ft) | Result |",
        "|---|---:|---:|---:|---|",
    ]
    for row in reconciliation:
        markdown.append(
            f"| {row['survey_date']} | {row['excel_nwp_ft']:.1f} | "
            f"{row['canonical_axial_ft']:.1f} | {row['difference_ft']:+.1f} | "
            f"{'PASS' if row['passes'] else 'FAIL'} |"
        )
    markdown.extend([
        "",
        "## Qualification",
        "",
        "`NO_GO_MODEL_READY_TRAJECTORY`",
        "",
        "The 2018 onset and follow-up observations still fail the frozen scalar reconciliation, "
        "and one section cannot support a generalisable forecasting claim. Phase-2 locking records "
        "the audited outcome; it does not authorise training.",
        "",
    ])
    (output / "phase2_finalization_report.md").write_text("\n".join(markdown))
    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
