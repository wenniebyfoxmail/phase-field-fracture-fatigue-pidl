#!/usr/bin/env python3
"""Build the frozen two-round blind TIFF measurement packet for S02-E004."""

from __future__ import annotations

import argparse
import binascii
import csv
from datetime import datetime, timedelta, timezone
import hashlib
import hmac
import json
import math
import os
import re
import secrets
import struct
import time
import urllib.request
import zlib
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image


SPECIMENS = ("H01-1", "H05-1", "V05-1")
ZIP_LOCKS = {
    "H01.zip": {
        "md5": "1a067a54b75d7f780ac0cb577aebcfcb",
        "tail_sha256": "bc13303f579536bd8bc08d8b1c805025dfe7fde5649392ec9480a4ddb7bdd959",
        "entries": 3686,
    },
    "H05.zip": {
        "md5": "8ec3a330f59a7178ab6eb838899321db",
        "tail_sha256": "6f2236acd67dc7029a58d6767a6217737bb6044aada1931296d78c817cfe5668",
        "entries": 3690,
    },
    "V05.zip": {
        "md5": "1dc696bd097e7e00ac092619163068ec",
        "tail_sha256": "699238b96f07302e4a8656ace07b18007173d5a56fbf9460d52be09145d6dc56",
        "entries": 3668,
    },
}
WORKBOOK_RECEIPT_SHA256 = "ca7d10beeb963c5629f9537f82bb6bbd487bb1d53806dafe831a5c55177b90fc"


def parse_zip64_extra(extra: bytes, usize: int, csize: int, offset: int):
    pos = 0
    while pos + 4 <= len(extra):
        header_id, length = struct.unpack_from("<HH", extra, pos)
        data = extra[pos + 4 : pos + 4 + length]
        pos += 4 + length
        if header_id != 1:
            continue
        cursor = 0
        if usize == 0xFFFFFFFF:
            usize = struct.unpack_from("<Q", data, cursor)[0]
            cursor += 8
        if csize == 0xFFFFFFFF:
            csize = struct.unpack_from("<Q", data, cursor)[0]
            cursor += 8
        if offset == 0xFFFFFFFF:
            offset = struct.unpack_from("<Q", data, cursor)[0]
        break
    return usize, csize, offset


def parse_central_directory(payload: bytes) -> list[dict]:
    entries, pos = [], 0
    while True:
        index = payload.find(b"PK\x01\x02", pos)
        if index < 0 or index + 46 > len(payload):
            break
        values = struct.unpack_from("<4s6H3I5H2I", payload, index)
        (_, _, _, flags, method, _, _, crc, csize, usize, name_len,
         extra_len, comment_len, _, _, _, offset) = values
        end = index + 46 + name_len + extra_len + comment_len
        if end > len(payload):
            break
        name_bytes = payload[index + 46 : index + 46 + name_len]
        encoding = "utf-8" if flags & 0x800 else "cp437"
        name = name_bytes.decode(encoding, errors="strict")
        extra = payload[index + 46 + name_len : index + 46 + name_len + extra_len]
        usize, csize, offset = parse_zip64_extra(extra, usize, csize, offset)
        entries.append({"name": name, "flags": flags, "method": method,
                        "crc32": crc, "compressed_size": csize,
                        "uncompressed_size": usize, "local_header_offset": offset})
        pos = end
    return entries


def http_range(url: str, start: int, end: int) -> bytes:
    expected = end - start + 1
    last_error = None
    for attempt in range(5):
        try:
            request = urllib.request.Request(url, headers={"Range": f"bytes={start}-{end}"})
            with urllib.request.urlopen(request, timeout=120) as response:
                data = response.read()
                if response.status != 206 or len(data) != expected:
                    raise RuntimeError(f"bad range response {response.status} {len(data)}/{expected}")
                content_range = response.headers.get("Content-Range", "")
                match = re.fullmatch(r"bytes (\d+)-(\d+)/(\d+|\*)", content_range)
                if not match or int(match.group(1)) != start or int(match.group(2)) != end:
                    raise RuntimeError(f"bad Content-Range: {content_range!r}")
                return data
        except Exception as error:
            last_error = error
            time.sleep(min(2**attempt, 8))
    raise RuntimeError(f"range request failed: {start}-{end}: {last_error}")


def remote_extract(url: str, entry: dict, target: Path) -> None:
    offset = int(entry["local_header_offset"])
    header = http_range(url, offset, offset + 29)
    signature, _, flags, method, _, _, _, _, _, name_len, extra_len = struct.unpack(
        "<4s5H3I2H", header
    )
    if (signature != b"PK\x03\x04" or flags != int(entry["flags"])
            or method != int(entry["method"]) or flags & 1 or method not in (0, 8)):
        raise RuntimeError(f"unsupported local member: {entry['name']}")
    local_name_bytes = http_range(url, offset + 30, offset + 30 + name_len - 1)
    encoding = "utf-8" if flags & 0x800 else "cp437"
    if local_name_bytes.decode(encoding, errors="strict") != entry["name"]:
        raise RuntimeError(f"local/central filename mismatch: {entry['name']}")
    start = offset + 30 + name_len + extra_len
    remaining = int(entry["compressed_size"])
    chunks, cursor = [], start
    while remaining:
        length = min(1024 * 1024, remaining)
        chunks.append(http_range(url, cursor, cursor + length - 1))
        cursor += length
        remaining -= length
    compressed = b"".join(chunks)
    payload = compressed if method == 0 else zlib.decompress(compressed, -15)
    if len(payload) != int(entry["uncompressed_size"]):
        raise RuntimeError(f"uncompressed size mismatch: {entry['name']}")
    if (binascii.crc32(payload) & 0xFFFFFFFF) != int(entry["crc32"]):
        raise RuntimeError(f"CRC-32 mismatch: {entry['name']}")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(payload)


def processed_rows(path: Path) -> list[tuple[int, float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(s for s in workbook.worksheets if "PROCESSED" in s.title.upper())
        rows = [(int(r[0]), float(r[1])) for r in sheet.iter_rows(
            min_row=3, max_col=2, values_only=True
        ) if r[0] is not None]
    finally:
        workbook.close()
    if len(rows) < 5 or any(b[0] <= a[0] for a, b in zip(rows, rows[1:])):
        raise ValueError(f"invalid processed rows: {path}")
    return rows


def state_indices(count: int) -> list[int]:
    indices = [round(k * (count - 1) / 4) for k in range(5)]
    if len(set(indices)) != 5:
        raise ValueError(f"nonunique state indices for count {count}: {indices}")
    return indices


def cycle_from_name(name: str) -> int:
    match = re.search(r"_(\d+)__\d+_0\.tif$", Path(name).name)
    if not match:
        raise ValueError(f"cannot parse camera-0 TIFF cycle: {name}")
    return int(match.group(1))


def local_half_interval(cycles: list[int], index: int) -> float:
    if index == 0:
        return float(cycles[1] - cycles[0])
    if index == len(cycles) - 1:
        return float(cycles[-1] - cycles[-2])
    return min(cycles[index] - cycles[index - 1], cycles[index + 1] - cycles[index]) / 2


def crc32_file(path: Path) -> int:
    checksum = 0
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            checksum = zlib.crc32(block, checksum)
    return checksum & 0xFFFFFFFF


def verify_workbook_receipt(root: Path) -> None:
    receipt_path = root / "local_archive/open_data/zenodo_10895431_preflight_20261008/metadata/selective_excel_extraction_receipt.json"
    if hashlib.sha256(receipt_path.read_bytes()).hexdigest() != WORKBOOK_RECEIPT_SHA256:
        raise ValueError("workbook receipt identity failure")
    members = json.loads(receipt_path.read_text(encoding="utf-8")).get("members")
    if not isinstance(members, list) or len(members) != 12:
        raise ValueError("workbook receipt must contain exactly twelve members")
    paths, identities = [], []
    expected_root = (root / "local_archive/open_data/zenodo_10895431_preflight_20261008/excel").resolve()
    for member in members:
        path = (root / member["target"]).resolve()
        if expected_root not in path.parents:
            raise ValueError(f"workbook target outside frozen root: {path}")
        if not path.is_file() or path.stat().st_size != int(member["uncompressed_size"]):
            raise ValueError(f"workbook identity size failure: {path}")
        if crc32_file(path) != int(member["crc32"]):
            raise ValueError(f"workbook identity CRC failure: {path}")
        paths.append(path); identities.append(path.stem.replace(" DATA", ""))
    if len(set(paths)) != 12 or len(set(identities)) != 12:
        raise ValueError("duplicate workbook path or identity")


def verify_tail(tail_path: Path, zip_name: str) -> list[dict]:
    payload = tail_path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    if len(payload) != 1024 * 1024 or digest != ZIP_LOCKS[zip_name]["tail_sha256"]:
        raise ValueError(f"tail identity failure: {zip_name}")
    entries = parse_central_directory(payload)
    if len(entries) != ZIP_LOCKS[zip_name]["entries"]:
        raise ValueError(f"central-directory entry-count failure: {zip_name}")
    if len({entry["name"] for entry in entries}) != len(entries):
        raise ValueError(f"duplicate central-directory member: {zip_name}")
    return entries


def opaque_token(secret: bytes, round_name: str, source_member: str) -> str:
    message = f"{round_name}\0{source_member}".encode()
    return hmac.new(secret, message, hashlib.sha256).hexdigest()[:16]


def sanitize_tiff(source: Path, destination: Path) -> tuple[int, int, str, str]:
    with Image.open(source) as image:
        if image.format != "TIFF":
            raise ValueError(f"not TIFF: {source}")
        image.load()
        width, height = image.size
        if width <= 0 or height <= 0:
            raise ValueError(f"invalid dimensions: {source}")
        mode = image.mode
        pixels = hashlib.sha256(image.tobytes()).hexdigest()
        clean = image.copy()
    destination.parent.mkdir(parents=True, exist_ok=True)
    clean.save(destination, format="TIFF", compression="tiff_lzw")
    with Image.open(destination) as check:
        check.verify()
    with Image.open(destination) as check:
        check.load()
        if check.size != (width, height) or check.mode != mode:
            raise ValueError(f"sanitized image geometry changed: {source}")
        if hashlib.sha256(check.tobytes()).hexdigest() != pixels:
            raise ValueError(f"sanitized image pixels changed: {source}")
        forbidden = {270, 271, 272, 306, 315, 33432}
        if forbidden.intersection(set(check.tag_v2.keys())):
            raise ValueError(f"identifying TIFF metadata survived: {destination}")
    os.utime(destination, (946684800, 946684800))
    return width, height, mode, pixels


def write_csv(path: Path, rows: list[dict], fields: list[str] | None = None) -> None:
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
        writer.writeheader(); writer.writerows(rows)


def annotation_fields() -> list[str]:
    return [
        "blind_id", "image_file", "visibility_status", "ruler_start_px",
        "ruler_end_px", "ruler_span_mm", "reference_plane_x_px", "tip_x_px",
        "projected_length_mm", "operator_notes",
    ]


def blank_annotation_row(blind_id: str, image_file: str) -> dict:
    row = {field: "" for field in annotation_fields()}
    row["blind_id"] = blind_id
    row["image_file"] = image_file
    return row


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def image_pixel_sha256(path: Path) -> str:
    with Image.open(path) as image:
        if image.format != "TIFF":
            raise ValueError(f"not TIFF: {path}")
        image.load()
        return hashlib.sha256(image.tobytes()).hexdigest()


def load_verified_selection(packet_dir: Path, sealed_dir: Path) -> list[dict]:
    receipt = json.loads((packet_dir / "packet_receipt.json").read_text(encoding="utf-8"))
    selection_path = sealed_dir / "selection.json"
    if file_sha256(selection_path) != receipt.get("selection_sha256"):
        raise ValueError("sealed selection digest mismatch")
    extraction_path = sealed_dir / "extraction_receipt.json"
    if file_sha256(extraction_path) != receipt.get("extraction_receipt_sha256"):
        raise ValueError("sealed extraction receipt digest mismatch")
    extraction = json.loads(extraction_path.read_text(encoding="utf-8"))
    if len(extraction) != 15 or len({row["source_member"] for row in extraction}) != 15:
        raise ValueError("sealed extraction receipt identity/cardinality failure")
    selected = json.loads(selection_path.read_text(encoding="utf-8"))
    if len(selected) != 15:
        raise ValueError("sealed selection must contain 15 frames")
    if len({item["source_member"] for item in selected}) != 15:
        raise ValueError("sealed selection has duplicate source members")
    if len({item["canonical_path"] for item in selected}) != 15:
        raise ValueError("sealed selection has duplicate canonical paths")
    for item in selected:
        canonical = sealed_dir / item["canonical_path"]
        if image_pixel_sha256(canonical) != item["pixel_sha256"]:
            raise ValueError(f"canonical pixel identity mismatch: {canonical}")
    return selected


def complete_round_a(packet_dir: Path, sealed_dir: Path) -> None:
    selected = load_verified_selection(packet_dir, sealed_dir)
    secret = (sealed_dir / "token_secret.bin").read_bytes()
    if len(secret) != 32:
        raise ValueError("sealed token secret identity failure")
    annotations_path = packet_dir / "round_A/annotations.csv"
    with annotations_path.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    expected = {opaque_token(secret, "A", item["source_member"]): item for item in selected}
    if len(rows) != 15 or {row["blind_id"] for row in rows} != set(expected):
        raise ValueError("Round A annotation identity/cardinality failure")
    round_images = list((packet_dir / "round_A").glob("*.tif"))
    expected_filenames = {f"m_{token}.tif" for token in expected}
    if len(round_images) != 15 or {path.name for path in round_images} != expected_filenames:
        raise ValueError("Round A image identity/cardinality failure")
    for token, item in expected.items():
        image_path = packet_dir / "round_A" / f"m_{token}.tif"
        if image_pixel_sha256(image_path) != item["pixel_sha256"]:
            raise ValueError(f"Round A image pixels changed: {token}")
        with Image.open(image_path) as image:
            image.load()
            if list(image.size) != [item["width_px"], item["height_px"]] or image.mode != item["image_mode"]:
                raise ValueError(f"Round A image geometry/mode changed: {token}")
    allowed = {"resolved", "ambiguous", "not_visible"}
    numeric_fields = (
        "ruler_start_px", "ruler_end_px", "ruler_span_mm",
        "reference_plane_x_px", "tip_x_px", "projected_length_mm",
    )
    for row in rows:
        if row["image_file"] != f"m_{row['blind_id']}.tif":
            raise ValueError(f"annotation image filename mismatch: {row['blind_id']}")
        status = row["visibility_status"].strip()
        if status not in allowed:
            raise ValueError(f"invalid visibility status: {row['blind_id']}")
        if status == "resolved":
            try:
                values = {field: float(row[field]) for field in numeric_fields}
            except (TypeError, ValueError) as error:
                raise ValueError(f"incomplete resolved measurement: {row['blind_id']}") from error
            if not all(math.isfinite(value) for value in values.values()):
                raise ValueError(f"nonfinite resolved measurement: {row['blind_id']}")
            width = float(expected[row["blind_id"]]["width_px"])
            for field in ("ruler_start_px", "ruler_end_px", "reference_plane_x_px", "tip_x_px"):
                if not 0.0 <= values[field] < width:
                    raise ValueError(f"coordinate outside image: {row['blind_id']} {field}")
            pixel_span = abs(values["ruler_end_px"] - values["ruler_start_px"])
            if pixel_span <= 0 or values["ruler_span_mm"] <= 0 or values["projected_length_mm"] < 0:
                raise ValueError(f"invalid ruler calibration: {row['blind_id']}")
            expected_length = (
                abs(values["tip_x_px"] - values["reference_plane_x_px"])
                * values["ruler_span_mm"] / pixel_span
            )
            if abs(expected_length - values["projected_length_mm"]) > 0.01:
                raise ValueError(f"projected length arithmetic mismatch: {row['blind_id']}")
        elif any(row[field].strip() for field in numeric_fields):
            raise ValueError(f"unresolved frame has numeric fields: {row['blind_id']}")
    completion_path = sealed_dir / "round_A_completion.json"
    if completion_path.exists():
        raise FileExistsError(completion_path)
    completion_path.write_text(json.dumps({
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "annotations_sha256": file_sha256(annotations_path),
        "row_count": 15,
    }, indent=2) + "\n", encoding="utf-8")
    os.chmod(completion_path, 0o600)
    print(json.dumps({"round": "A", "status": "COMPLETE", "record": str(completion_path)}, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--round", choices=("A", "complete-A", "B"), required=True)
    parser.add_argument("--project-root", type=Path, required=True)
    parser.add_argument("--tail-dir", type=Path, required=True)
    parser.add_argument("--packet-dir", type=Path, required=True)
    parser.add_argument("--sealed-dir", type=Path, required=True)
    args = parser.parse_args()
    root = args.project_root.resolve()
    packet_dir = args.packet_dir.resolve()
    sealed_dir = args.sealed_dir.resolve()
    if packet_dir == sealed_dir or packet_dir in sealed_dir.parents or sealed_dir in packet_dir.parents:
        raise ValueError("packet and sealed directories must be separate trees")
    metadata_path = root / "local_archive/open_data/zenodo_10895431_preflight_20261008/metadata/zenodo_record_10895431_retrieved_20261008.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    files = {item["key"]: item for item in metadata["files"]}
    for zip_name, lock in ZIP_LOCKS.items():
        if files[zip_name]["checksum"] != f"md5:{lock['md5']}":
            raise ValueError(f"metadata ZIP identity failure: {zip_name}")
    verify_workbook_receipt(root)

    if args.round == "complete-A":
        complete_round_a(packet_dir, sealed_dir)
        return

    if args.round == "B":
        round_dir = packet_dir / "round_B"
        if round_dir.exists():
            raise FileExistsError(round_dir)
        completion_path = sealed_dir / "round_A_completion.json"
        completion = json.loads(completion_path.read_text(encoding="utf-8"))
        completed = datetime.fromisoformat(completion["completed_at"])
        if completed.tzinfo is None or datetime.now(timezone.utc) < completed.astimezone(timezone.utc) + timedelta(hours=24):
            raise ValueError("Round B embargo has not elapsed")
        annotations_path = packet_dir / "round_A/annotations.csv"
        if file_sha256(annotations_path) != completion.get("annotations_sha256"):
            raise ValueError("Round A annotations changed after completion")
        selected = load_verified_selection(packet_dir, sealed_dir)
        secret = (sealed_dir / "token_secret.bin").read_bytes()
        if len(secret) != 32:
            raise ValueError("sealed token secret identity failure")
        round_dir.mkdir(parents=True)
        rows, tokens = [], []
        a_tokens = {opaque_token(secret, "A", item["source_member"]) for item in selected}
        for item in selected:
            token = opaque_token(secret, "B", item["source_member"])
            tokens.append(token)
            source = sealed_dir / item["canonical_path"]
            destination = round_dir / f"m_{token}.tif"
            sanitize_tiff(source, destination)
            rows.append(blank_annotation_row(token, destination.name))
        if len(set(tokens)) != 15 or set(tokens).intersection(a_tokens):
            raise ValueError("Round B token uniqueness/blinding failure")
        rows.sort(key=lambda row: row["blind_id"])
        write_csv(round_dir / "annotations.csv", rows, annotation_fields())
        (round_dir / "ROUND_B_RELEASE.json").write_text(json.dumps({
            "round_a_completed_at": completion["completed_at"],
            "round_b_released_at": datetime.now(timezone.utc).isoformat(),
            "minimum_delay_hours": 24,
        }, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"round": "B", "tasks": 15, "output": str(round_dir)}, indent=2))
        return

    if packet_dir.exists() or sealed_dir.exists():
        raise FileExistsError("Round A packet or sealed directory already exists")
    packet_dir.mkdir(parents=True)
    sealed_dir.mkdir(parents=True, mode=0o700)
    selected = []
    extraction_receipt = []
    for specimen in SPECIMENS:
        group = specimen.split("-")[0]
        workbook = root / f"local_archive/open_data/zenodo_10895431_preflight_20261008/excel/{group}/{specimen} DATA.xlsx"
        rows = processed_rows(workbook)
        entries = [e for e in verify_tail(
            args.tail_dir / f"{group}.zip.tail1m", f"{group}.zip"
        ) if e["name"].startswith(specimen + "/") and e["name"].endswith("_0.tif")]
        entries.sort(key=lambda e: cycle_from_name(e["name"]))
        image_cycles = [cycle_from_name(e["name"]) for e in entries]
        if len(image_cycles) != len(set(image_cycles)):
            raise ValueError(f"duplicate camera-0 cycles: {specimen}")
        for ordinal, row_index in enumerate(state_indices(len(rows))):
            processed_cycle, workbook_a = rows[row_index]
            choices = sorted(
                enumerate(image_cycles), key=lambda pair: (abs(pair[1] - processed_cycle), pair[1])
            )
            image_index, image_cycle = choices[0]
            difference = abs(image_cycle - processed_cycle)
            tolerance = local_half_interval(image_cycles, image_index)
            if difference > tolerance:
                raise ValueError(
                    f"mapping outside half interval: {specimen} {processed_cycle} -> {image_cycle}"
                )
            entry = entries[image_index]
            selected.append({
                "specimen": specimen, "group": group, "state_ordinal": ordinal,
                "processed_row_index": row_index, "processed_cycle": processed_cycle,
                "workbook_a_m": workbook_a, "tiff_cycle": image_cycle,
                "cycle_difference": image_cycle - processed_cycle,
                "mapping_half_interval": tolerance, "source_member": entry["name"],
                "source_zip": f"{group}.zip", "crc32": f"{entry['crc32']:08x}",
                "uncompressed_size": entry["uncompressed_size"], "entry": entry,
            })
    if len(selected) != 15:
        raise AssertionError(f"expected 15 selected frames, got {len(selected)}")
    if len({item["source_member"] for item in selected}) != 15:
        raise ValueError("selected TIFF members are not unique")

    canonical_dir = sealed_dir / "canonical_images"
    canonical_dir.mkdir()
    for number, item in enumerate(selected, start=1):
        raw_target = sealed_dir / f"raw_{number:03d}.tif"
        remote_extract(files[item["source_zip"]]["links"]["self"], item["entry"], raw_target)
        extraction_receipt.append({
            "source_zip": item["source_zip"],
            "source_zip_md5": ZIP_LOCKS[item["source_zip"]]["md5"],
            "tail_sha256": ZIP_LOCKS[item["source_zip"]]["tail_sha256"],
            "source_url": files[item["source_zip"]]["links"]["self"],
            "source_member": item["source_member"],
            "flags": item["entry"]["flags"],
            "method": item["entry"]["method"],
            "crc32": f"{item['entry']['crc32']:08x}",
            "compressed_size": item["entry"]["compressed_size"],
            "uncompressed_size": item["entry"]["uncompressed_size"],
            "local_header_offset": item["entry"]["local_header_offset"],
            "processed_cycle": item["processed_cycle"],
            "tiff_cycle": item["tiff_cycle"],
            "cycle_difference": item["cycle_difference"],
        })
        canonical = canonical_dir / f"s{number:03d}.tif"
        width, height, mode, pixel_sha = sanitize_tiff(raw_target, canonical)
        raw_target.unlink()
        item["width_px"], item["height_px"], item["image_mode"] = width, height, mode
        item["pixel_sha256"] = pixel_sha
        item["canonical_path"] = str(canonical.relative_to(sealed_dir))
        del item["entry"]
    selection_path = sealed_dir / "selection.json"
    selection_path.write_text(json.dumps(selected, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.chmod(selection_path, 0o600)
    extraction_path = sealed_dir / "extraction_receipt.json"
    extraction_path.write_text(json.dumps(extraction_receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.chmod(extraction_path, 0o600)
    secret_path = sealed_dir / "token_secret.bin"
    secret_path.write_bytes(secrets.token_bytes(32))
    os.chmod(secret_path, 0o600)
    secret = secret_path.read_bytes()

    round_dir = packet_dir / "round_A"
    round_dir.mkdir()
    rows, tokens = [], []
    for item in selected:
        token = opaque_token(secret, "A", item["source_member"])
        tokens.append(token)
        destination = round_dir / f"m_{token}.tif"
        sanitize_tiff(sealed_dir / item["canonical_path"], destination)
        rows.append(blank_annotation_row(token, destination.name))
    b_tokens = {opaque_token(secret, "B", item["source_member"]) for item in selected}
    if len(set(tokens)) != 15 or len(b_tokens) != 15 or set(tokens).intersection(b_tokens):
        raise ValueError("opaque token uniqueness/blinding failure")
    rows.sort(key=lambda row: row["blind_id"])
    write_csv(round_dir / "annotations.csv", rows, annotation_fields())
    if len(list(round_dir.glob("*.tif"))) != 15 or len(rows) != 15:
        raise ValueError("Round A packet cardinality failure")
    (packet_dir / "packet_receipt.json").write_text(json.dumps({
        "experiment_id": "S02-E004", "protocol_revision": "v1",
        "unique_frames": 15, "released_rounds": ["A"], "tasks_in_round_A": 15,
        "round_B_status": "NOT_CREATED__24H_AFTER_ROUND_A_REQUIRED",
        "selection_sha256": file_sha256(selection_path),
        "extraction_receipt_sha256": file_sha256(extraction_path),
        "source_zip_md5": {k: files[k]["checksum"] for k in ("H01.zip", "H05.zip", "V05.zip")},
    }, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"round": "A", "selected_frames": 15, "tasks": 15,
                      "packet_dir": str(packet_dir), "sealed_dir": str(sealed_dir)}, indent=2))


if __name__ == "__main__":
    main()
