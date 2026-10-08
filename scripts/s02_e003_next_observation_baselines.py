#!/usr/bin/env python3
"""Evaluate frozen persistence and secant baselines for S02-E003.

The evaluator is read-only with respect to the source workbooks. It does not
fit parameters, pool specimens, clip predictions, or use future-assisted
derived columns.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import zlib
from pathlib import Path
from statistics import mean, median

from openpyxl import load_workbook


EXPECTED_SPECIMENS = {
    "H01-1", "H01-2", "H01-3", "H01-4",
    "H05-1", "H05-2", "H05-3",
    "V05-1", "V05-2", "V05-3", "V05-4",
}


def analytic_self_test() -> None:
    cases = {
        "constant": ([0.0, 1.0, 3.0, 6.0], [2.0, 2.0, 2.0, 2.0], 2.0),
        "linear": ([0.0, 1.0, 3.0, 6.0], [1.0, 3.0, 7.0, 13.0], 13.0),
        "accelerating": ([0.0, 1.0, 2.0, 3.0], [0.0, 1.0, 3.0, 6.0], 5.0),
    }
    for name, (cycles, cracks, expected) in cases.items():
        pred = cracks[2] + (cracks[2] - cracks[1]) / (cycles[2] - cycles[1]) * (
            cycles[3] - cycles[2]
        )
        if not math.isclose(pred, expected, rel_tol=0.0, abs_tol=1e-12):
            raise AssertionError(f"analytic fixture failed: {name}: {pred} != {expected}")


def processed_rows(path: Path) -> tuple[list[float], list[float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = next(
            (s for s in workbook.worksheets if "PROCESSED" in s.title.upper()),
            None,
        )
        if sheet is None:
            raise ValueError(f"missing processed sheet: {path.name}")
        cycles: list[float] = []
        cracks: list[float] = []
        raw_rows = list(sheet.iter_rows(min_row=3, max_col=2, values_only=True))
        last_data_index = max(
            (i for i, row in enumerate(raw_rows) if row[0] is not None or row[1] is not None),
            default=-1,
        )
        for row_number, row in enumerate(raw_rows[: last_data_index + 1], start=3):
            if row[0] is None and row[1] is None:
                raise ValueError(f"internal blank row: {path.name} row {row_number}")
            if row[0] is None or row[1] is None:
                raise ValueError(f"missing required cell: {path.name} row {row_number}")
            cycle, crack = float(row[0]), float(row[1])
            if not math.isfinite(cycle) or not math.isfinite(crack):
                raise ValueError(f"nonfinite required cell: {path.name} row {row_number}")
            cycles.append(cycle)
            cracks.append(crack)
        return cycles, cracks
    finally:
        workbook.close()


def validate_trajectory(specimen: str, cycles: list[float], cracks: list[float]) -> None:
    if len(cycles) != len(cracks) or len(cycles) < 4:
        raise ValueError(f"invalid observation count: {specimen}")
    if not all(b > a for a, b in zip(cycles, cycles[1:])):
        raise ValueError(f"cycles not strictly increasing: {specimen}")
    if not all(b >= a for a, b in zip(cracks, cracks[1:])):
        raise ValueError(f"crack length decreases: {specimen}")


def crc32_file(path: Path) -> int:
    checksum = 0
    with path.open("rb") as stream:
        while block := stream.read(1024 * 1024):
            checksum = zlib.crc32(block, checksum)
    return checksum & 0xFFFFFFFF


def verify_receipt(input_root: Path, receipt_path: Path) -> list[Path]:
    payload = json.loads(receipt_path.read_text(encoding="utf-8"))
    members = payload.get("members")
    if not isinstance(members, list) or len(members) != 12:
        raise ValueError("receipt must contain exactly twelve members")
    project_root = receipt_path.resolve().parents[4]
    expected_paths: list[Path] = []
    names: list[str] = []
    for member in members:
        if member.get("crc_verified") is not True:
            raise ValueError(f"receipt member was not CRC verified: {member.get('name')}")
        path = (project_root / str(member["target"])).resolve()
        if path.parent.parent != input_root.resolve():
            raise ValueError(f"receipt target outside frozen input root: {path}")
        if not path.is_file():
            raise ValueError(f"missing receipt target: {path}")
        if path.stat().st_size != int(member["uncompressed_size"]):
            raise ValueError(f"size mismatch: {path.name}")
        if crc32_file(path) != int(member["crc32"]):
            raise ValueError(f"CRC-32 mismatch: {path.name}")
        expected_paths.append(path)
        names.append(path.stem.replace(" DATA", ""))
    if len(set(expected_paths)) != 12 or len(set(names)) != 12:
        raise ValueError("receipt contains duplicate paths or specimen identities")
    actual_paths = sorted(path.resolve() for path in input_root.glob("*/*.xlsx"))
    if sorted(expected_paths) != actual_paths:
        raise ValueError("input root does not exactly match the twelve receipt targets")
    return sorted(expected_paths)


def quantile(values: list[float], probability: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return ordered[lower]
    fraction = position - lower
    return ordered[lower] * (1.0 - fraction) + ordered[upper] * fraction


def evaluate(input_root: Path, receipt_path: Path) -> tuple[list[dict], list[dict], dict]:
    files = verify_receipt(input_root, receipt_path)
    observed = {path.stem.replace(" DATA", "") for path in files}
    if "H05-4" not in observed:
        raise ValueError("frozen H05-4 workbook is absent")
    usable = observed - {"H05-4"}
    if usable != EXPECTED_SPECIMENS:
        raise ValueError(
            f"frozen usable specimen mismatch: missing={sorted(EXPECTED_SPECIMENS-usable)} "
            f"extra={sorted(usable-EXPECTED_SPECIMENS)}"
        )

    predictions: list[dict] = []
    specimen_rows: list[dict] = []
    for path in files:
        specimen = path.stem.replace(" DATA", "")
        if specimen == "H05-4":
            continue
        group = specimen.split("-")[0]
        cycles, cracks = processed_rows(path)
        validate_trajectory(specimen, cycles, cracks)
        local: list[dict] = []
        for i in range(2, len(cycles) - 1):
            dt_prev = cycles[i] - cycles[i - 1]
            dt_next = cycles[i + 1] - cycles[i]
            persistence = cracks[i]
            slope = (cracks[i] - cracks[i - 1]) / dt_prev
            secant = cracks[i] + slope * dt_next
            target = cracks[i + 1]
            row = {
                "group": group,
                "specimen": specimen,
                "origin_index": i,
                "origin_cycle": cycles[i],
                "target_cycle": cycles[i + 1],
                "horizon_cycles": dt_next,
                "origin_crack_m": cracks[i],
                "target_crack_m": target,
                "persistence_prediction_m": persistence,
                "secant_prediction_m": secant,
                "persistence_error_m": persistence - target,
                "secant_error_m": secant - target,
                "persistence_absolute_error_m": abs(persistence - target),
                "secant_absolute_error_m": abs(secant - target),
            }
            if not all(math.isfinite(float(v)) for v in row.values() if isinstance(v, (int, float))):
                raise ValueError(f"nonfinite prediction: {specimen} origin {i}")
            local.append(row)
            predictions.append(row)
        p_mae = mean(float(r["persistence_absolute_error_m"]) for r in local)
        s_mae = mean(float(r["secant_absolute_error_m"]) for r in local)
        specimen_rows.append({
            "group": group,
            "specimen": specimen,
            "observation_count": len(cycles),
            "eligible_origin_count": len(local),
            "persistence_mae_m": p_mae,
            "secant_mae_m": s_mae,
            "secant_minus_persistence_mae_m": s_mae - p_mae,
            "persistence_bias_m": mean(float(r["persistence_error_m"]) for r in local),
            "secant_bias_m": mean(float(r["secant_error_m"]) for r in local),
            "secant_wins": s_mae < p_mae,
        })

    p_macro = mean(float(r["persistence_mae_m"]) for r in specimen_rows)
    s_macro = mean(float(r["secant_mae_m"]) for r in specimen_rows)
    group_summary = {}
    for group in sorted({str(r["group"]) for r in specimen_rows}):
        selected = [r for r in specimen_rows if r["group"] == group]
        group_summary[group] = {
            "specimen_count": len(selected),
            "persistence_macro_mae_m": mean(float(r["persistence_mae_m"]) for r in selected),
            "secant_macro_mae_m": mean(float(r["secant_mae_m"]) for r in selected),
        }
    horizons = [float(r["horizon_cycles"]) for r in predictions]
    persistence_absolute_errors = [
        float(r["persistence_absolute_error_m"]) for r in predictions
    ]
    secant_absolute_errors = [float(r["secant_absolute_error_m"]) for r in predictions]
    summary = {
        "experiment_id": "S02-E003",
        "protocol_revision": "v1",
        "run_id": "S02-E003-R001",
        "validity": "PASS",
        "specimen_count": len(specimen_rows),
        "eligible_origin_count": len(predictions),
        "persistence_macro_mae_m": p_macro,
        "secant_macro_mae_m": s_macro,
        "relative_macro_mae_change": (
            (s_macro - p_macro) / p_macro if p_macro != 0.0 else None
        ),
        "secant_specimen_wins": sum(bool(r["secant_wins"]) for r in specimen_rows),
        "persistence_macro_bias_m": mean(float(r["persistence_bias_m"]) for r in specimen_rows),
        "secant_macro_bias_m": mean(float(r["secant_bias_m"]) for r in specimen_rows),
        "persistence_absolute_error_p50_m": quantile(persistence_absolute_errors, 0.50),
        "persistence_absolute_error_p90_m": quantile(persistence_absolute_errors, 0.90),
        "secant_absolute_error_p50_m": quantile(secant_absolute_errors, 0.50),
        "secant_absolute_error_p90_m": quantile(secant_absolute_errors, 0.90),
        "primary_pass": s_macro < p_macro,
        "scientific_verdict": "supports" if s_macro < p_macro else "negative",
        "horizon_cycles_min": min(horizons),
        "horizon_cycles_median": median(horizons),
        "horizon_cycles_max": max(horizons),
        "group_diagnostics": group_summary,
        "blocked_claims": [
            "calibrated observation uncertainty",
            "image-based crack measurement validity",
            "independent ground truth",
            "fixed-cycle forecasting",
            "road or material generalisation",
        ],
    }
    return predictions, specimen_rows, summary


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-root", type=Path, required=True)
    parser.add_argument("--receipt", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    analytic_self_test()
    predictions, specimens, summary = evaluate(args.input_root, args.receipt)
    args.output_dir.mkdir(parents=True, exist_ok=False)
    write_csv(args.output_dir / "predictions.csv", predictions)
    write_csv(args.output_dir / "specimen_metrics.csv", specimens)
    (args.output_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    print(json.dumps(summary, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
