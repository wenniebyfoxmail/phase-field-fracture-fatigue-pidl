#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path


DEFAULT_FEM_WORKBOOK = Path(
    "/Users/wenxiaofang/Library/CloudStorage/OneDrive-UniversityofCambridge/"
    "PIDL result/_pidl_handoff_fem_u012_soft_hist0_whole_field_package_20260603/"
    "nstep2/fem22_nstep2_c0_c5_region_summary.xlsx"
)
DEFAULT_DATA_DIR = (
    Path(__file__).resolve().parents[1]
    / "_analysis_fem_mechanism_20260528"
    / "allfemmesh_lambda_hist_20260607"
)


METHOD_STYLE = {
    "FEM nstep2 latest": {"color": "#000000", "marker": "o", "linestyle": "-", "linewidth": 2.0},
    "PIDL strict FEM mesh": {"color": "#0072B2", "marker": "s", "linestyle": "-", "linewidth": 1.7},
    "PIDL strict FEM mesh + lambda/Eirres": {
        "color": "#D55E00",
        "marker": "^",
        "linestyle": "-",
        "linewidth": 1.7,
    },
}


def _to_float(value: object) -> float:
    if value is None or value == "":
        return math.nan
    return float(value)


def load_fem_peak_energy(workbook: Path) -> list[dict[str, float | str]]:
    from openpyxl import load_workbook

    wb = load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["states"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    idx = {key: i for i, key in enumerate(header)}
    out = []
    seen_cycles = set()
    for row in rows[1:]:
        cycle = int(row[idx["cycle"]])
        load_factor = _to_float(row[idx["load_factor"]])
        if cycle <= 0 or cycle in seen_cycles:
            continue
        if abs(load_factor - 1.0) > 1e-4:
            continue
        if not bool(row[idx["after_history_refresh"]]):
            continue
        out.append(
            {
                "method": "FEM nstep2 latest",
                "physical_cycle": cycle,
                "source_cycle": cycle,
                "state_label": str(row[idx["state_label"]]),
                "E_el": _to_float(row[idx["E_el"]]),
                "E_d_degraded": _to_float(row[idx["E_d"]]),
                "E_d_raw": math.nan,
                "E_irres": math.nan,
                "E_total_degraded": _to_float(row[idx["total_energy"]]),
                "E_total_raw": math.nan,
            }
        )
        seen_cycles.add(cycle)
    return out


def load_pidl_energy(path: Path, method: str, max_physical_cycle: int) -> list[dict[str, float | str]]:
    out = []
    with path.open(newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            source_cycle = int(row["cycle"])
            physical_cycle = source_cycle + 1
            if physical_cycle > max_physical_cycle:
                break
            out.append(
                {
                    "method": method,
                    "physical_cycle": physical_cycle,
                    "source_cycle": source_cycle,
                    "state_label": f"trained_1NN_{source_cycle}.pt",
                    "E_el": _to_float(row["E_el"]),
                    "E_d_degraded": _to_float(row.get("E_d_degraded", row.get("E_d"))),
                    "E_d_raw": _to_float(row.get("E_d_raw")),
                    "E_irres": _to_float(row.get("E_irres")),
                    "E_total_degraded": _to_float(row.get("E_total_degraded", row.get("E_total"))),
                    "E_total_raw": _to_float(row.get("E_total_raw")),
                }
            )
    return out


def write_csv(path: Path, rows: list[dict[str, float | str]]) -> None:
    fieldnames = [
        "method",
        "physical_cycle",
        "source_cycle",
        "state_label",
        "E_el",
        "E_d_degraded",
        "E_d_raw",
        "E_irres",
        "E_total_degraded",
        "E_total_raw",
    ]
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def read_csv(path: Path) -> list[dict[str, float | str]]:
    rows: list[dict[str, float | str]] = []
    with path.open(newline="") as handle:
        for row in csv.DictReader(handle):
            rows.append(row)
    return rows


def write_summary(path: Path, rows: list[dict[str, float | str]]) -> None:
    methods = []
    for row in rows:
        method = str(row["method"])
        if method not in methods:
            methods.append(method)
    quantities = ["E_el", "E_d_degraded", "E_d_raw", "E_irres", "E_total_degraded", "E_total_raw"]
    with path.open("w", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["method", "quantity", "cycle1", "cycle5", "delta_c5_minus_c1", "min", "max"],
        )
        writer.writeheader()
        for method in methods:
            subset = [row for row in rows if row["method"] == method]
            for quantity in quantities:
                values = [float(row[quantity]) for row in subset if not math.isnan(float(row[quantity]))]
                if not values:
                    continue
                writer.writerow(
                    {
                        "method": method,
                        "quantity": quantity,
                        "cycle1": f"{values[0]:.16e}",
                        "cycle5": f"{values[-1]:.16e}",
                        "delta_c5_minus_c1": f"{values[-1] - values[0]:.16e}",
                        "min": f"{min(values):.16e}",
                        "max": f"{max(values):.16e}",
                    }
                )


def _series(rows: list[dict[str, float | str]], method: str, quantity: str) -> tuple[list[float], list[float]]:
    subset = [row for row in rows if row["method"] == method]
    cycles = [float(row["physical_cycle"]) for row in subset]
    values = [float(row[quantity]) for row in subset]
    return cycles, values


def plot_components(rows: list[dict[str, float | str]], out_png: Path, out_pdf: Path) -> None:
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    panels = [
        ("E_el", r"$E_{\mathrm{el}}$", "Elastic energy"),
        ("E_d_degraded", r"$E_d$ effective", "Damage energy"),
        ("E_total_degraded", r"$E_{\mathrm{total}}$ effective", "Total energy"),
    ]
    methods = list(METHOD_STYLE)
    fig, axes = plt.subplots(1, 3, figsize=(12.8, 3.8), sharex=True)
    for ax, (quantity, ylabel, title) in zip(axes, panels):
        for method in methods:
            cycles, values = _series(rows, method, quantity)
            style = METHOD_STYLE[method]
            ax.plot(cycles, values, label=method, markersize=4.8, **style)
        ax.set_title(title, fontsize=10.5)
        ax.set_xlabel("Physical cycle")
        ax.set_ylabel(ylabel)
        ax.grid(True, color="0.88", linewidth=0.7)
        ax.set_xticks([1, 2, 3, 4, 5])
    fig.suptitle("Latest FEM nstep2 energy versus strict FEM-mesh PIDL, Umax=0.12", fontsize=12)
    handles = [
        Line2D(
            [0],
            [0],
            color=METHOD_STYLE[method]["color"],
            marker=METHOD_STYLE[method]["marker"],
            linestyle=METHOD_STYLE[method]["linestyle"],
            linewidth=METHOD_STYLE[method]["linewidth"],
            label=method,
        )
        for method in methods
    ]
    fig.legend(handles=handles, loc="lower center", ncol=3, frameon=False, fontsize=9)
    fig.subplots_adjust(left=0.065, right=0.99, top=0.78, bottom=0.25, wspace=0.34)
    fig.savefig(out_png, dpi=300)
    fig.savefig(out_pdf)
    plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, default=DEFAULT_DATA_DIR)
    parser.add_argument("--fem-workbook", type=Path, default=DEFAULT_FEM_WORKBOOK)
    parser.add_argument("--no-plot", action="store_true")
    parser.add_argument("--plot-existing", action="store_true")
    args = parser.parse_args()

    out_csv = args.data_dir / "latest_fem_nstep2_vs_allfemmesh_pidl_energy_c1_c5.csv"
    out_summary = args.data_dir / "latest_fem_nstep2_vs_allfemmesh_pidl_energy_summary.csv"
    out_png = args.data_dir / "latest_fem_nstep2_vs_allfemmesh_pidl_energy_c1_c5.png"
    out_pdf = args.data_dir / "latest_fem_nstep2_vs_allfemmesh_pidl_energy_c1_c5.pdf"

    if args.plot_existing:
        rows = read_csv(out_csv)
    else:
        fem_rows = load_fem_peak_energy(args.fem_workbook)
        max_physical_cycle = max(int(row["physical_cycle"]) for row in fem_rows)
        baseline_rows = load_pidl_energy(
            args.data_dir / "baseline_energy_components.csv",
            "PIDL strict FEM mesh",
            max_physical_cycle,
        )
        adapthist_rows = load_pidl_energy(
            args.data_dir / "adapthist_energy_components.csv",
            "PIDL strict FEM mesh + lambda/Eirres",
            max_physical_cycle,
        )
        rows = fem_rows + baseline_rows + adapthist_rows
        write_csv(out_csv, rows)
        write_summary(out_summary, rows)
    if not args.no_plot:
        plot_components(rows, out_png, out_pdf)
    print(out_csv)
    print(out_summary)
    if not args.no_plot:
        print(out_png)
        print(out_pdf)


if __name__ == "__main__":
    main()
