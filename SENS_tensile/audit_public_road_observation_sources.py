#!/usr/bin/env python3
"""Audit small public road-observation payloads without fusing unrelated sites."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re
from typing import Any
from zipfile import ZipFile


SOURCE_URLS = {
    "ltpp_distress": "https://s3.amazonaws.com/zip.aims.infopave.com/SDR29/AIMS/DIS/MDS_MDP/06/DIS_MDS_MDP_06B410.zip",
    "ltpp_profile": "https://s3.amazonaws.com/zip.aims.infopave.com/SDR29/AIMS/PRF/LPF_ERD/06/PRF_LPF_ERD_06B410.zip",
    "ltpp_maintenance": "https://s3.amazonaws.com/zip.aims.infopave.com/SDR29/AIMS/MNT/06/MNT_06B410.zip",
    "ltpp_manual": "https://infopave.fhwa.dot.gov/InfoPave_Repository/files/IMS-171.pdf",
    "mnroad_sensors": "https://www.dot.state.mn.us/mnroad/instrumentation/sensor_locations.xlsx",
    "mnroad_traffic": "https://www.dot.state.mn.us/mnroad/files/mnroad-mainline-traffic-calculator-1994-07-2013-09-updated.xlsx",
    "mnroad_schedule": "https://www.dot.state.mn.us/mnroad/data/files/MnROAD%20Monitoring%20Schedule.pdf",
    "mnroad_weather": "https://www.dot.state.mn.us/mnroad/files/mnroad-weather-2014.pdf",
    "idics_description": "https://drive.google.com/uc?id=1vEFb6IbAhsq80PB0gbba3ozD6263aEv8",
    "idics_sample1": "https://drive.google.com/uc?id=1UUL4-mzJhp0Ffl1E0e6gcbXdGSP_VXUh",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_erd_header(payload: bytes) -> dict[str, Any]:
    """Parse the stable header subset used by LTPP ERD profile files."""
    text = payload.decode("latin-1", errors="replace")
    header = text.split("\nEND", 1)[0]

    def value(prefix: str) -> str:
        for line in header.splitlines():
            if line.startswith(prefix):
                return line[len(prefix) :].strip()
        return ""

    first_data = next(
        (line for line in header.splitlines()[1:] if re.match(r"\s*\d+\s*,", line)),
        "",
    )
    numbers = [token.strip() for token in first_data.split(",") if token.strip()]
    return {
        "title": value("TITLE"),
        "survey_date": value("DATE"),
        "x_units": value("XUNITS"),
        "sample_count": int(numbers[1]) if len(numbers) > 1 else None,
        "sample_spacing": float(numbers[5]) if len(numbers) > 5 else None,
        "profile_units": value("UNITSNAM"),
        "section": value(" Section No:"),
        "subsection": value(" SubSection of:"),
    }


def inspect_ltpp(root: Path) -> dict[str, Any]:
    distress_path = root / "DIS_MDS_MDP_06B410.zip"
    profile_path = root / "PRF_LPF_ERD_06B410.zip"
    maintenance_path = root / "MNT_06B410.zip"
    for path in (distress_path, profile_path, maintenance_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    with ZipFile(distress_path) as archive:
        distress_members = [name for name in archive.namelist() if name.lower().endswith(".jpg")]
    distress_dates = sorted({name.split("/", 1)[0] for name in distress_members})

    profiles: list[dict[str, Any]] = []
    with ZipFile(profile_path) as archive:
        for name in archive.namelist():
            if name.lower().endswith(".erd"):
                profiles.append({"member": name, **parse_erd_header(archive.read(name))})
    profile_dates = sorted({row["member"].split("/", 1)[0] for row in profiles})
    sample_counts = sorted({row["sample_count"] for row in profiles})
    sample_spacings = sorted({row["sample_spacing"] for row in profiles})

    with ZipFile(maintenance_path) as archive:
        maintenance_members = archive.namelist()

    return {
        "distress_file_count": len(distress_members),
        "distress_dates": distress_dates,
        "profile_file_count": len(profiles),
        "profile_dates": profile_dates,
        "profile_sample_counts": sample_counts,
        "profile_sample_spacings_m": sample_spacings,
        "maintenance_file_count": len(maintenance_members),
        # The source is a handwritten scan. These dates are deliberately labelled
        # as manual visual transcription rather than machine-readable records.
        "maintenance_dates_manual_visual": ["1993-08-01", "1996-09-01", "2000-05-01"],
        "maintenance_work_codes": "not_machine_verified",
    }


def inspect_mnroad(root: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific guidance
        raise RuntimeError("openpyxl is required to inspect MnROAD workbooks") from exc

    sensor_path = root / "sensor_locations.xlsx"
    traffic_path = root / "mainline_traffic_calculator_1994_2013.xlsx"
    for path in (sensor_path, traffic_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    sensor_sheet = load_workbook(sensor_path, read_only=True, data_only=True).active
    sensor_rows = list(sensor_sheet.iter_rows(min_row=2, values_only=True))
    sensor_models: dict[str, int] = {}
    cell_ids: set[str] = set()
    cell22_models: dict[str, int] = {}
    cell22_active_2009_models: dict[str, int] = {}
    for row in sensor_rows:
        cell, model = row[0], str(row[1])
        cell_ids.add(str(cell))
        sensor_models[model] = sensor_models.get(model, 0) + 1
        if cell == 22 and model in {"LE", "TE", "PG", "TC"}:
            cell22_models[model] = cell22_models.get(model, 0) + 1
            installed, removed = row[4], row[5]
            if installed <= datetime(2009, 1, 1) and (removed is None or removed >= datetime(2009, 1, 1)):
                cell22_active_2009_models[model] = cell22_active_2009_models.get(model, 0) + 1

    traffic_book = load_workbook(traffic_path, read_only=True, data_only=True)
    traffic_rows = 0
    dates: list[datetime] = []
    for sheet_name in ("Lane1", "Lane2"):
        sheet = traffic_book[sheet_name]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if isinstance(row[0], datetime) and row[1] in {1, 2}:
                dates.append(row[0])
                traffic_rows += 1

    return {
        "sensor_rows": len(sensor_rows),
        "sensor_cells": len(cell_ids),
        "sensor_top_models": dict(sorted(sensor_models.items(), key=lambda item: (-item[1], item[0]))[:12]),
        "cell22_selected_models": cell22_models,
        "cell22_active_2009_selected_models": cell22_active_2009_models,
        "traffic_daily_lane_rows": traffic_rows,
        "traffic_date_start": min(dates).date().isoformat(),
        "traffic_date_end": max(dates).date().isoformat(),
        "traffic_channels": "FHWA classes C1-C15|total volume|heavy commercial|BESAL|CESAL",
    }


def inspect_idics(root: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    try:
        import numpy as np
        from PIL import Image
        from skimage.registration import phase_cross_correlation
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment-specific guidance
        raise RuntimeError("numpy, Pillow, scikit-image and openpyxl are required") from exc

    description_path = root / "2D_Challenge_1.0_set_descriptions.xlsx"
    sample_path = root / "Sample1.zip"
    for path in (description_path, sample_path):
        if not path.is_file():
            raise FileNotFoundError(path)

    sheet = load_workbook(description_path, read_only=True, data_only=True)["Overview"]
    sample1 = next(row for row in sheet.iter_rows(min_row=2, values_only=True) if row[1] == "Sample1")
    known_increment = 0.05
    with ZipFile(sample_path) as archive:
        names = sorted(name for name in archive.namelist() if name.lower().endswith(".tif"))

        def load(index: int) -> Any:
            return np.asarray(
                Image.open(BytesIO(archive.read(f"trxy_s2_{index:02d}.tif"))),
                dtype=np.float32,
            )

        reference = load(0)
        checks = []
        for index in (1, 2, 5, 10, 20):
            shift, error, _ = phase_cross_correlation(
                reference,
                load(index),
                upsample_factor=100,
                normalization=None,
            )
            measured = float((abs(float(shift[0])) + abs(float(shift[1]))) / 2)
            expected = known_increment * index
            checks.append(
                {
                    "frame": index,
                    "expected_abs_shift_px_each_axis": expected,
                    "measured_abs_shift_px_mean_axes": measured,
                    "absolute_error_px": abs(measured - expected),
                    "registration_error": float(error),
                }
            )

    return (
        {
            "image_count": len(names),
            "image_shape": f"{reference.shape[0]}x{reference.shape[1]}",
            "method": sample1[2],
            "contrast": str(sample1[3]),
            "declared_noise_gray_levels": sample1[4],
            "declared_increment_px_each_axis": known_increment,
            "max_registration_absolute_error_px": max(row["absolute_error_px"] for row in checks),
        },
        checks,
    )


def classify_source_role(source_id: str) -> tuple[str, str]:
    roles = {
        "ltpp_06B410": (
            "section-level real deterioration trajectory",
            "not crack-tip mechanics or full-field hidden-state truth",
        ),
        "mnroad": (
            "load/environment/mechanical observation bridge after same-cell numeric export",
            "metadata and WIM aggregates alone are not state assimilation",
        ),
        "pavetrack": (
            "visual crack morphology and registration pretraining",
            "not mechanical validation or physical RUL without scale and co-located channels",
        ),
        "idics_2d": (
            "DIC observation-operator metrology and noise calibration",
            "not road-domain training or fracture-mechanism evidence",
        ),
    }
    return roles[source_id]


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        raise ValueError(f"cannot write empty table: {path}")
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build_tables(ltpp: dict[str, Any], mnroad: dict[str, Any], idics: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    source_rows = []
    source_specs = [
        ("ltpp_06B410", "real road section", datetime.strptime(ltpp["profile_dates"][0], "%Y%m%d").date().isoformat(), datetime.strptime(ltpp["distress_dates"][-1], "%Y%m%d").date().isoformat(), "06B410 section and dated survey directories"),
        ("mnroad", "real instrumented pavement facility", mnroad["traffic_date_start"], mnroad["traffic_date_end"], "cell/station/offset/depth sensor registry"),
        ("pavetrack", "real repeated crack imagery", "2022-04-18", "2023-12-31", "location-isolated IDs; no public route coordinates or physical scale"),
        ("idics_2d", "synthetic/experimental DIC benchmark", "not_calendar_time", "not_calendar_time", "image coordinates with declared subpixel motion"),
    ]
    for source_id, evidence, start, end, registration in source_specs:
        allowed, prohibited = classify_source_role(source_id)
        source_rows.append(
            {
                "source_id": source_id,
                "evidence_class": evidence,
                "date_start": start,
                "date_end": end,
                "spatial_registration": registration,
                "allowed_use": allowed,
                "prohibited_use": prohibited,
                "audit_status": "downloaded_and_checked",
            }
        )

    channel_rows = [
        {"channel": "crack/distress image", "ltpp": "direct_6_images_one_date", "mnroad": "available_not_in_payload", "pavetrack": "direct_8928_pairs", "idics": "not_applicable", "current_use": "visual morphology only"},
        {"channel": "longitudinal profile", "ltpp": "direct_32_runs_four_dates", "mnroad": "method_documented_not_in_payload", "pavetrack": "missing", "idics": "missing", "current_use": "section condition/roughness trajectory"},
        {"channel": "maintenance", "ltpp": "scan_three_manual_dates", "mnroad": "not_in_payload", "pavetrack": "missing", "idics": "not_applicable", "current_use": "reset/event covariate after code verification"},
        {"channel": "traffic exposure", "ltpp": "available_via_InfoPave_not_in_payload", "mnroad": "direct_daily_C1_C15_BESAL_CESAL", "pavetrack": "missing", "idics": "not_applicable", "current_use": "calendar-to-load exposure for MnROAD mainline"},
        {"channel": "temperature/environment", "ltpp": "available_via_InfoPave_not_in_payload", "mnroad": "registry_and_protocol_only_numeric_share_unavailable", "pavetrack": "missing", "idics": "not_applicable", "current_use": "request/alignment design only"},
        {"channel": "strain/displacement/pressure", "ltpp": "missing", "mnroad": "sensor_registry_only_numeric_share_unavailable", "pavetrack": "missing", "idics": "subpixel_observation_operator_benchmark", "current_use": "DIC metrology now; MnROAD assimilation after export"},
        {"channel": "FWD basin", "ltpp": "InfoPave_table_export_not_in_payload", "mnroad": "protocol_documented_numeric_export_not_in_payload", "pavetrack": "missing", "idics": "not_applicable", "current_use": "blocked until same-section numeric export"},
        {"channel": "FEM latent fields", "ltpp": "not_observed", "mnroad": "not_observed", "pavetrack": "not_observed", "idics": "not_observed", "current_use": "never label as sensor data"},
    ]

    alignment_rows = [
        {"candidate_packet": "LTPP_06B410", "observed_timeline": "profile 1997-04-11/1998-12-15/1999-11-03/2000-06-25; distress 2000-06-28; maintenance 1993/1996/2000 manual scan", "same_asset": "yes_section_level", "mechanical_closure": "no", "decision": "usable for section-level condition/reset study; not mechanism assimilation"},
        {"candidate_packet": "MnROAD_Cell22_2009_2013", "observed_timeline": "daily mainline WIM aggregate plus registered LE/TE/PG/TC sensor locations; numeric sensor/FWD/distress data still requested", "same_asset": "designed_but_incomplete", "mechanical_closure": "potential_after_export", "decision": "highest-priority real mechanical packet"},
        {"candidate_packet": "PaveTrack_location_trajectory", "observed_timeline": "repeated images and masks within public location ID", "same_asset": "yes_image_space", "mechanical_closure": "no", "decision": "visual observation operator only"},
        {"candidate_packet": "iDICs_Sample1", "observed_timeline": f"21 frames; declared 0.05 px/frame each axis; max audit error {idics['max_registration_absolute_error_px']:.3f} px", "same_asset": "yes_synthetic_image", "mechanical_closure": "known_kinematic_truth_only", "decision": "calibrate DIC registration/noise; never pool as road samples"},
    ]
    return source_rows, channel_rows, alignment_rows


def raw_inventory(raw_root: Path) -> list[dict[str, Any]]:
    rows = []
    by_name = {
        "DIS_MDS_MDP_06B410.zip": "ltpp_distress",
        "PRF_LPF_ERD_06B410.zip": "ltpp_profile",
        "MNT_06B410.zip": "ltpp_maintenance",
        "IMS-171.pdf": "ltpp_manual",
        "sensor_locations.xlsx": "mnroad_sensors",
        "mainline_traffic_calculator_1994_2013.xlsx": "mnroad_traffic",
        "MnROAD_Monitoring_Schedule.pdf": "mnroad_schedule",
        "mnroad-weather-2014.pdf": "mnroad_weather",
        "2D_Challenge_1.0_set_descriptions.xlsx": "idics_description",
        "Sample1.zip": "idics_sample1",
    }
    for path in sorted(raw_root.rglob("*")):
        if not path.is_file():
            continue
        key = by_name.get(path.name, "unregistered")
        rows.append(
            {
                "relative_path": path.relative_to(raw_root).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
                "source_url": SOURCE_URLS.get(key, ""),
            }
        )
    return rows


def decision_text(ltpp: dict[str, Any], mnroad: dict[str, Any], idics: dict[str, Any]) -> str:
    return f"""# Public road-observation source audit

Date: 2026-07-31

## Decision

The downloaded sources are useful, but they belong to different evidence layers
and must not be concatenated as if they were one road trajectory.

- **LTPP 06B410** is now a real section-level pilot: {ltpp['profile_file_count']}
  longitudinal-profile runs cover four survey dates from 1997 to 2000, six
  distress images were taken on 2000-06-28, and the maintenance scan contains
  three manually readable dates. This can test condition trajectories and
  maintenance resets. It cannot validate crack-tip active support because no
  co-registered full-field mechanics exists.
- **MnROAD** is the strongest path to real mechanical assimilation. The local
  payload contains {mnroad['sensor_rows']:,} located sensor records across
  {mnroad['sensor_cells']} cells and {mnroad['traffic_daily_lane_rows']:,} daily
  lane records from {mnroad['traffic_date_start']} to {mnroad['traffic_date_end']}
  with FHWA class counts and BESAL/CESAL. It does not yet contain numeric
  strain, pressure, temperature, FWD or distress time series. The public data
  share timed out, so those channels remain requested rather than evaluated.
- **PaveTrack** remains a real visual observation source only. Its 8,928
  location-image-mask pairs are useful for morphology, segmentation and
  registration, but lack physical scale and paired mechanical/load channels.
- **iDICs 2D Sample1** is a metrology control, not road data. Its {idics['image_count']}
  512x512 images have a declared 0.05-pixel shift per frame per axis; the local
  subpixel check recovered the selected shifts with at most
  {idics['max_registration_absolute_error_px']:.3f} px absolute error.

## Correct integration

1. Use iDICs to qualify the image-to-displacement/strain observation operator.
2. Use PaveTrack to qualify visual crack morphology and missing-registration
   handling across held-out locations.
3. Use LTPP for section-level calendar deterioration and maintenance-reset
   prediction, with profile/distress/load/climate channels joined only by the
   same section and date.
4. Request one complete MnROAD same-cell packet, beginning with Cell 22 over
   2009-2013: LE/TE strain, PG pressure, TC temperature, FWD drops, distress
   surveys and mainline WIM/ESAL. This is the first candidate that can connect
   loading, environment, mechanical response and visible damage.
5. Evaluate each task separately. Cross-source pretraining is allowed; sample-
   level fusion across LTPP, MnROAD, PaveTrack and iDICs is prohibited.

## Claim boundary

This audit upgrades public-source availability from unknown to checked. It does
not set `real_mechanical_road_data_evaluated=true`: the required same-road
numeric mechanical packet is still absent. FEM damage, history, raw driver,
degradation and active driver remain latent references, never direct sensors.
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--raw-root", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    ltpp = inspect_ltpp(args.raw_root / "ltpp_06B410")
    mnroad = inspect_mnroad(args.raw_root / "mnroad")
    idics, dic_checks = inspect_idics(args.raw_root / "idics")
    sources, channels, alignments = build_tables(ltpp, mnroad, idics)

    write_csv(args.out / "raw_asset_inventory.csv", raw_inventory(args.raw_root))
    write_csv(args.out / "source_inventory.csv", sources)
    write_csv(args.out / "channel_availability.csv", channels)
    write_csv(args.out / "temporal_alignment.csv", alignments)
    write_csv(args.out / "dic_registration_check.csv", dic_checks)
    (args.out / "decision.md").write_text(decision_text(ltpp, mnroad, idics), encoding="utf-8")

    manifest = {
        "schema_version": "public_road_observation_audit_v1",
        "created_at": datetime.now().astimezone().isoformat(),
        "raw_root": str(args.raw_root.resolve()),
        "raw_assets": raw_inventory(args.raw_root),
        "ltpp_summary": ltpp,
        "mnroad_summary": mnroad,
        "idics_summary": idics,
        "pavetrack_packages": [
            "analysis/pavetrack_real_observation_pilot_20260731",
            "analysis/pavetrack_registered_observation_gate_20260731",
        ],
        "decision": {
            "public_sources_downloaded_and_audited": True,
            "real_visual_data_evaluated": True,
            "real_mechanical_road_data_evaluated": False,
            "sample_level_cross_source_fusion_allowed": False,
            "next_same_asset_packet": "MnROAD Cell 22 2009-2013",
        },
    }
    (args.out / "RUN_MANIFEST.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    assets = sorted(path for path in args.out.iterdir() if path.name != "HASHES.sha256")
    (args.out / "HASHES.sha256").write_text(
        "".join(f"{sha256(path)}  {path.name}\n" for path in assets if path.is_file()),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
