"""Read-only trajectory inventory. No predictions, fitting, model loads or interpolation."""
import csv
from datetime import datetime,timezone
import hashlib
import io
import json
import math
from pathlib import Path
import re
import signal
import socket
import os
import subprocess
import time
import warnings
import zipfile
import openpyxl

ROOT=Path('/Users/wenxiaofang/phase-field-fracture-with-pidl')
RAW=ROOT/'local_archive/open_data/history_candidate_selection_20261017/raw'
MN=Path('/Users/wenxiaofang/Downloads/MnROAD/HMA Distress All export.xlsx')
ME=ROOT/'local_archive/open_data/mendeley_four_specimen_audit_20261013/4. Cycles vs Crack Length.xlsx'
MANIFEST=ROOT/'local_archive/open_data/mendeley_four_specimen_audit_20261013/archive_header_manifest.json'
PHM=RAW/'PHM2019_crack_growth_training_data_v2.zip'


def check(ok,msg):
    if not ok:raise ValueError(msg)


def numeric(v):
    return isinstance(v,(int,float)) and not isinstance(v,bool) and math.isfinite(v)


def write_csv(p,rs):
    with p.open('w',newline='') as f:
        wr=csv.DictWriter(f,fieldnames=list(rs[0]),lineterminator='\n');wr.writeheader();wr.writerows(rs)


def exact_windows(keys,h):
    s=set(keys)
    return sum(all(k+j*h in s for j in [-2,-1,1]) for k in s)


def mendeley():
    w=openpyxl.load_workbook(ME,read_only=True,data_only=True)
    tables=[];unions={};manifest=json.loads(MANIFEST.read_text())
    mg={}
    for r in manifest:
        match=re.search(r'/([^/]+)/[^/]+/[^/]+_(\d+)\.csv$',r['name'])
        check(match is not None,'unknown Mendeley member schema')
        mg.setdefault(match[1],set()).add(int(match[2]))
    for sh in w:
        vals=list(sh.values);check(vals[0][:3]==('Figure ID','Cycles','Crack Length (Ground Truth)'),'Mendeley label header')
        rs=[r for r in vals[1:] if r[0] is not None]
        check(all(numeric(r[0]) and numeric(r[1]) and numeric(r[2]) and int(r[1])==30*(int(r[0])-1) for r in rs),'Mendeley indexing')
        base=sh.title.split('(')[0]
        group={'E_0_175_1':'E_0Deg_17.5%_#1','C_15_20_1':'C_15Deg_20%_#1'}[base]
        check(all(int(r[0]) in mg[group] for r in rs),'Mendeley label-member stage missing')
        keys=[int(r[1]) for r in rs];check(keys==sorted(set(keys)),'Mendeley duplicate/unordered keys')
        tables.append({'sheet':sh.title,'specimen_code':base,'label_rows':len(keys),'first_cycle':min(keys),'last_cycle':max(keys),
                       **{f'four_state_windows_h{h}':exact_windows(keys,h) for h in [30,600,1200]}})
        unions.setdefault(base,set()).update(keys)
    summary={'sheets':tables,'sheet_count':len(tables),'distinct_label_codes':len(unions),
        'archive_group_count':len(mg),'archive_csv_count':len(manifest),'union_is_conditional_not_author_confirmed':True,
        'unions':{k:{'unique_label_cycles':len(v),**{f'four_state_windows_h{h}':exact_windows(v,h) for h in [30,600,1200]}} for k,v in unions.items()},
        'independence_phase_reference_label_error':'unresolved; see prior Mendeley qualification',
        'predicted_column_used':False}
    return tables,summary


def mnroad():
    w=openpyxl.load_workbook(MN,read_only=True,data_only=True)
    rows=iter(w['Export Worksheet'].values);headers=next(rows)
    rs=[dict(zip(headers,r)) for r in rows];groups={}
    for r in rs:groups.setdefault(r['SECTION_ID'],[]).append(r)
    fields=['LONG_NWP_L_L','LONG_NWP_L_M','LONG_NWP_L_H']
    pattern=re.compile(r'seal|patch|microsurfac|reconstruct|overlay|rehab|mill|repair',re.I)
    tab=[];annual=[]
    for sid,rr in groups.items():
        rr.sort(key=lambda r:(r['DAY'],r['ID']))
        variant_count=len({(r['CELL'],r['LANE'],r['SECTION_START_DATE'],r['SECTION_END_DATE']) for r in rr})
        dates=[r['DAY'] for r in rr];slots={}
        for r in rr:
            if 3<=r['DAY'].month<=5:slots[r['DAY'].year]=r
        # Last observation within the declared spring slot; fail any incomplete selected label.
        complete={yr:r for yr,r in slots.items() if all(numeric(r[f]) for f in fields)}
        identity_clear=variant_count==1 and len(dates)==len(set(dates))
        origins=[yr for yr in sorted(complete) if all(yr+j in complete for j in [-2,-1,1])] if identity_clear else []
        flags=[r for r in rr if pattern.search(str(r['COMMENTS'] or ''))]
        tab.append({'section_id':sid,'cell_code':rr[0]['CELL'],'lane':rr[0]['LANE'],'state_rows':len(rr),
            'distinct_dates':len(set(dates)),'duplicate_date_rows':len(rr)-len(set(dates)),
            'metadata_variant_count':variant_count,'all_cell_codes':json.dumps(sorted({r['CELL'] for r in rr})),
            'identity_clear_for_inventory':identity_clear,
            'section_start':rr[0]['SECTION_START_DATE'].date().isoformat(),'section_end':rr[0]['SECTION_END_DATE'].date().isoformat(),
            'dates_outside_section':sum(not(r['SECTION_START_DATE']<=r['DAY']<=r['SECTION_END_DATE']) for r in rr),
            'incomplete_nwp_rows':sum(not all(numeric(r[f]) for f in fields) for r in rr),
            'maintenance_comment_warning_rows':len(flags),'complete_spring_slots':len(complete),
            'four_spring_slot_origins_upper_bound':len(origins),
            'old_cell22_excluded_from_new_shortlist':str(rr[0]['CELL'])=='22',
            'intervention_clearance':'not established'})
        for yr in origins:
            dates4=[complete[yr+j]['DAY'] for j in [-2,-1,0,1]]
            annual.append({'section_id':sid,'cell_code':rr[0]['CELL'],'origin_year':yr,
                'past2_date':dates4[0].date().isoformat(),'past1_date':dates4[1].date().isoformat(),
                'origin_date':dates4[2].date().isoformat(),'target_date':dates4[3].date().isoformat(),
                'actual_future_gap_days_not_input':(dates4[3]-dates4[2]).days,
                'origin_slot_cutoff':f'{yr}-05-31','target_slot_cutoff':f'{yr+1}-05-31',
                'maintenance_cleared':False,'fresh_cell_candidate':str(rr[0]['CELL'])!='22'})
    cell={r['CELL'] for r in rs}
    candidates=[r for r in tab if r['four_spring_slot_origins_upper_bound'] and not r['old_cell22_excluded_from_new_shortlist']]
    summary={'raw_rows':len(rs),'columns':len(headers),'section_ids':len(groups),'distinct_cell_codes_not_certified_independence':len(cell),
        'lanes':dict(__import__('collections').Counter(r['LANE'] for r in rs)),
        'sections_at_least_4_dates':sum(r['distinct_dates']>=4 for r in tab),
        'sections_with_duplicate_dates':sum(r['duplicate_date_rows']>0 for r in tab),
        'metadata_conflict_sections':[r for r in tab if r['metadata_variant_count']>1],
        'identity_clear_section_count':sum(r['identity_clear_for_inventory'] for r in tab),
        'incomplete_nwp_rows':sum(r['incomplete_nwp_rows'] for r in tab),
        'warning_rows':sum(r['maintenance_comment_warning_rows'] for r in tab),
        'spring_coverage_sections_excluding_cell22':len(candidates),
        'spring_coverage_cells_excluding_cell22':len({r['cell_code'] for r in candidates}),
        'spring_origins_upper_bound_excluding_cell22':sum(r['four_spring_slot_origins_upper_bound'] for r in candidates),
        'qualified_independent_intervention_cleared_cells':0,
        'zero_qualified_means_not_yet_qualified_not_proven_unusable':True,
        'shortlist_by_coverage_not_target_values':sorted(candidates,key=lambda r:(-r['four_spring_slot_origins_upper_bound'],r['section_id']))[:12],
        'sql':w['SQL']['A2'].value,'forecasting_fixed_365_day_claim':False}
    return tab,annual,summary


def phm():
    warnings.filterwarnings('ignore',message='Unknown extension is not supported')
    z=zipfile.ZipFile(PHM);check(z.testzip() is None,'PHM CRC failure')
    tab=[];joins=[]
    for s in [f'T{i}' for i in range(1,7)]:
        name=next(n for n in z.namelist() if n.endswith('/Description_'+s+'.xlsx'))
        w=openpyxl.load_workbook(io.BytesIO(z.read(name)),read_only=True,data_only=True)
        vals=[(int(r[0]),float(r[1])) for r in w['Sheet1'].values if numeric(r[0]) and numeric(r[1])]
        labelkeys=[r[0] for r in vals];check(len(labelkeys)==len(set(labelkeys)),'duplicate PHM labels')
        signals={};aliases=[];time_warnings=[];invalid_waves=[]
        for n in z.namelist():
            match=re.search('/'+s+r'/(\d+)/(signal|singal)_([12])\.csv$',n)
            if not match:continue
            if match[2]=='singal':aliases.append(n)
            cycle=int(match[1]);rep=int(match[3])
            check(rep not in signals.setdefault(cycle,{}),'duplicate PHM repeat')
            rows=list(csv.reader(io.StringIO(z.read(n).decode(),newline='')))
            check(rows[0]==['time','ch1','ch2'],'PHM signal header')
            vv=[];invalid=0
            for raw in rows[1:]:
                if not raw:continue
                try:
                    values=[float(x) for x in raw]
                    valid=len(values)==3 and all(math.isfinite(v) for v in values)
                except ValueError:
                    valid=False
                vv.append(values if valid else None)
                invalid+=not valid
            if invalid:invalid_waves.append({'member':n,'invalid_rows':invalid})
            adjacent=[(a,b) for a,b in zip(vv,vv[1:]) if a is not None and b is not None]
            equal=sum(b[0]==a[0] for a,b in adjacent)
            backward=sum(b[0]<a[0] for a,b in adjacent)
            if equal or backward:time_warnings.append({'member':n,'duplicate_time_steps':equal,'backward_time_steps':backward})
            signals[cycle][rep]=len(vv)
        check(all(set(v)=={1,2} for v in signals.values()),'PHM repeats missing')
        exact=sorted(set(labelkeys)&set(signals))
        bad=[b<=a for a,b in zip(labelkeys,labelkeys[1:])]
        tab.append({'specimen':s,'label_rows':len(vals),'signal_cycles':len(signals),'exact_joined_cycles':len(exact),
            'signal_only_cycles':json.dumps(sorted(set(signals)-set(labelkeys))),
            'label_only_cycles':json.dumps(sorted(set(labelkeys)-set(signals))),
            'source_label_order_monotonic':not any(bad),'typo_spelling_signal_members':len(aliases),
            'waveform_files_with_invalid_rows':len(invalid_waves),
            'waveform_invalid_rows':sum(r['invalid_rows'] for r in invalid_waves),
            'waveform_files_with_time_warnings':len(time_warnings),
            'waveform_duplicate_time_steps':sum(r['duplicate_time_steps'] for r in time_warnings),
            'waveform_backward_time_steps':sum(r['backward_time_steps'] for r in time_warnings),
            'three_past_plus_next_count_upper_bound':max(len(exact)-3,0),
            'fixed_h1000_four_state_windows':exact_windows(exact,1000),
            'typo_cycle_repaired':False})
        for cycle in sorted(set(labelkeys)|set(signals)):
            joins.append({'specimen':s,'cycle':cycle,'has_optical_label':cycle in labelkeys,'has_two_waveforms':cycle in signals,
                'signal_1_samples':signals.get(cycle,{}).get(1),'signal_2_samples':signals.get(cycle,{}).get(2)})
    return tab,joins,{'zip_bytes':PHM.stat().st_size,'zip_sha256':hashlib.sha256(PHM.read_bytes()).hexdigest(),
        'zip_members':len(z.namelist()),'crc_all_members_pass':True,'specimens':tab,
        'label_rows_total':sum(r['label_rows'] for r in tab),'signal_cycles_total':sum(r['signal_cycles'] for r in tab),
        'exact_joined_total':sum(r['exact_joined_cycles'] for r in tab),
        'three_past_plus_next_upper_bound':sum(r['three_past_plus_next_count_upper_bound'] for r in tab),
        'fixed_h1000_windows_total':sum(r['fixed_h1000_four_state_windows'] for r in tab),
        'tables_are_development_exposed':True,'T7_T8_accessed':False,
        'optical_error_quantified':False,'fixed_horizon_training_ready':False}


def main():
    out=ROOT/'local_archive/experiments/S02-E006/runs/S02-E006-R004/audit'
    check(not out.exists(),'fresh output required')
    check(not subprocess.check_output(['git','status','--porcelain'],text=True).strip(),'clean code snapshot required')
    head=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()
    auth=json.loads((out.parent/'authorization.json').read_text())
    check(auth['code_commit']==head and auth['review']=='CODE_REVIEW_PASS' and auth['training'] is False,'review/authority')
    start=time.monotonic();started=datetime.now(timezone.utc).isoformat()
    def timeout(*_):raise TimeoutError('audit60s cap')
    signal.signal(signal.SIGALRM,timeout);signal.alarm(60)
    original={p:p.read_bytes() for p in [MN,ME,MANIFEST,PHM]}
    mt,ms=mendeley();nt,na,ns=mnroad();pt,pj,ps=phm()
    out.mkdir(parents=True)
    for name,rs in [('mendeley_sheets.csv',mt),('mnroad_sections.csv',nt),('mnroad_spring_slot_inventory.csv',na),('phm_specimens.csv',pt),('phm_cycle_join.csv',pj)]:write_csv(out/name,rs)
    summary={'run_id':'S02-E006-R004','scope':'schema identity availability only, no forecasts','Mendeley':ms,'MnROAD':ns,'PHM2019':ps,'training':False,'fitting':False,'scientific_prediction_result':False}
    (out/'summary.json').write_text(json.dumps(summary,indent=2,default=str)+'\n')
    check(all(p.read_bytes()==b for p,b in original.items()),'source changed')
    receipt={'run_id':'S02-E006-R004','code_commit':head,'dirty_status':'clean at launch','hostname':socket.gethostname(),'pid':os.getpid(),'start_utc':started,'elapsed_seconds':time.monotonic()-start,'out':str(out),'input_paths':[str(p) for p in original],'input_bytes_unchanged':True,'execution_status':'succeeded','retrieval':'local verified','training':False,'fitting':False,'model_forward':False,'command':__import__('sys').argv}
    (out/'execution_receipt.json').write_text(json.dumps(receipt,indent=2)+'\n');signal.alarm(0)
    print(json.dumps(summary,indent=2,default=str))

if __name__=='__main__':main()
